#!/usr/bin/env python3
"""診断.xlsx から式場データ (js/data/venues.js) を生成する。

使い方:
    python3 tools/build_data.py source/診断.xlsx

Excel の解釈:
    シート1: A列=式場名、2行目以降が式場。B列以降の1行目がカテゴリ名。
             セルに文字が入っていれば、その式場はその列のカテゴリに該当する。
    シート2: A列2行目以降がカテゴリ一覧（診断に使うカテゴリ）。

式場名・カテゴリ名は Excel の文字列をそのまま使う（表記の修正・統合はしない）。
想定と異なる構造を見つけた場合は、生成を中止してエラーを表示する。
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "js" / "data" / "venues.js"
MARK_BEGIN = "/*__DATA_BEGIN__*/"
MARK_END = "/*__DATA_END__*/"


def fail(msg):
    print("エラー: " + msg, file=sys.stderr)
    sys.exit(1)


def main():
    if len(sys.argv) < 2:
        fail("Excelファイルのパスを指定してください")
    src = Path(sys.argv[1])
    wb = openpyxl.load_workbook(src, data_only=True)
    if len(wb.worksheets) < 2:
        fail("シートが2枚未満です")
    ws_v, ws_c = wb.worksheets[0], wb.worksheets[1]

    header = [c.value for c in ws_v[1]]
    if header[0] != "式場リスト":
        fail("シート1のA1が「式場リスト」ではありません: %r" % header[0])
    col_cats = header[1:]
    if any(not isinstance(h, str) or not h for h in col_cats):
        fail("シート1の1行目に空のカテゴリ列があります")

    sheet2 = [r[0] for r in ws_c.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    if ws_c["A1"].value != "カテゴリ":
        fail("シート2のA1が「カテゴリ」ではありません")

    warnings = []
    if col_cats != sheet2:
        only1 = [c for c in col_cats if c not in sheet2]
        only2 = [c for c in sheet2 if c not in col_cats]
        fail("シート1の列見出しとシート2のカテゴリが一致しません。シート1のみ=%s シート2のみ=%s" % (only1, only2))

    venues = []
    seen = set()
    for row in ws_v.iter_rows(min_row=2):
        name = row[0].value
        cells = row[1:len(col_cats) + 1]
        if name is None:
            if any(c.value is not None for c in cells):
                fail("%d行目: 式場名が空なのにカテゴリが入力されています" % row[0].row)
            continue
        if name in seen:
            fail("式場名が重複しています: %s" % name)
        seen.add(name)
        cats = []
        for cat, cell in zip(col_cats, cells):
            v = cell.value
            if v is None or (isinstance(v, str) and v.strip() == ""):
                continue
            if v != cat:
                # 想定: セルの文字は列見出しと同じ。異なる場合も「文字が入っている＝該当」と扱い、警告に残す。
                warnings.append("%s: 列「%s」のセル値が見出しと異なります (%r)" % (cell.coordinate, cat, v))
            cats.append(cat)
        if not cats:
            warnings.append("%d行目「%s」: カテゴリが1つも設定されていません" % (row[0].row, name))
        venues.append({"id": "v%02d" % (len(venues) + 1), "row": row[0].row, "name": name, "categories": cats})

    data = {
        "source": {"file": src.name, "venueSheet": ws_v.title, "categorySheet": ws_c.title},
        "categories": sheet2,
        "venues": venues,
        "warnings": warnings,
    }
    body = json.dumps(data, ensure_ascii=False, indent=1)
    OUT.write_text(
        "/* 自動生成ファイル: tools/build_data.py により %s から生成。直接編集しないでください。\n"
        "   式場やカテゴリを変更する場合は Excel を修正し、スクリプトを再実行してください。 */\n"
        "(function (root) {\n"
        "  root.VENUE_DIAG = root.VENUE_DIAG || {};\n"
        "  root.VENUE_DIAG.DATA = %s%s%s;\n"
        "})(typeof window !== 'undefined' ? window : globalThis);\n" % (src.name, MARK_BEGIN, body, MARK_END),
        encoding="utf-8",
    )
    print("式場 %d件 / カテゴリ %d件 を書き出しました: %s" % (len(venues), len(sheet2), OUT))
    for w in warnings:
        print("警告: " + w)


if __name__ == "__main__":
    main()
