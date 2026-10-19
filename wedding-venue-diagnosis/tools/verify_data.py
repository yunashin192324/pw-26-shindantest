#!/usr/bin/env python3
"""js/data/venues.js が 診断.xlsx と完全に一致しているかを、Excel を読み直して照合する。

使い方:
    python3 tools/verify_data.py source/診断.xlsx
生成スクリプトとは別の読み方（セルを座標で直接読む）で照合し、1件でも差があれば終了コード1で終わる。
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent


def load_js_data():
    text = (ROOT / "js" / "data" / "venues.js").read_text(encoding="utf-8")
    body = text.split("/*__DATA_BEGIN__*/", 1)[1].split("/*__DATA_END__*/", 1)[0]
    return json.loads(body)


def main():
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    ws = wb["シート1"]
    ws2 = wb["シート2"]
    js = load_js_data()
    problems = []

    cats_sheet2 = []
    r = 2
    while ws2.cell(r, 1).value is not None:
        cats_sheet2.append(ws2.cell(r, 1).value)
        r += 1
    if js["categories"] != cats_sheet2:
        problems.append("カテゴリ一覧がシート2と一致しません")

    max_col = ws.max_column
    headers = {c: ws.cell(1, c).value for c in range(2, max_col + 1)}
    excel = {}
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if name is None:
            continue
        excel[name] = (r, [headers[c] for c in range(2, max_col + 1) if ws.cell(r, c).value not in (None, "")])

    js_map = {v["name"]: v for v in js["venues"]}
    if len(js_map) != len(js["venues"]):
        problems.append("JS側で式場名が重複しています")
    for name in excel:
        if name not in js_map:
            problems.append("JSに無い式場: " + name)
    for name, v in js_map.items():
        if name not in excel:
            problems.append("Excelに無い式場: " + name)
            continue
        row, cats = excel[name]
        if v["row"] != row:
            problems.append("%s: 行番号の不一致 (%s != %s)" % (name, v["row"], row))
        if v["categories"] != cats:
            problems.append("%s: カテゴリ不一致 JS=%s Excel=%s" % (name, v["categories"], cats))

    print("Excel: 式場 %d件 / カテゴリ %d件" % (len(excel), len(cats_sheet2)))
    print("JS   : 式場 %d件 / カテゴリ %d件" % (len(js["venues"]), len(js["categories"])))
    print("付与されたカテゴリの総数: Excel %d / JS %d" % (
        sum(len(c) for _, c in excel.values()), sum(len(v["categories"]) for v in js["venues"])))
    if problems:
        print("不一致あり:")
        for p in problems:
            print(" - " + p)
        sys.exit(1)
    print("OK: 全式場・全カテゴリが一致しました")


if __name__ == "__main__":
    main()
